#!/usr/bin/env python3
"""
dateien.py - Dateien finden und ihren Text lesen.

Beantwortet zwei Fragen: Welche Dateien in einem Ordner kommen fuer die
Suche in Frage (dateien_im_ordner), und was steht in einer Datei drin
(lies_datei, in_abschnitte_teilen). Weiss nichts vom Index und nichts vom
Modell.
"""

import os

from smartsearch.kern import ocr

UNTERSTUETZT = (".txt", ".md", ".pdf", ".docx", ".xlsx", ".pptx")

# Ordnernamen, die NIE mitindexiert werden sollen, egal wo sie im
# durchsuchten Ordnerbaum auftauchen. Das sind App-eigene/Bibliotheks-
# Ordner (venv, Build-Ausgaben, Python-Cache, Git-Interna) - deren
# enthaltene Dateien (z.B. Vorlagen-.docx/.pptx aus installierten
# Bibliotheken) haben mit den eigentlichen Nutzer-Dokumenten nichts zu
# tun, blähen aber den Index unnötig auf und kosten Rechenzeit.
IGNORIERTE_ORDNERNAMEN = {"venv", ".venv", ".git", "__pycache__", "build", "dist", "node_modules"}


def lies_datei(pfad):
    ext = os.path.splitext(pfad)[1].lower()
    try:
        if ext in [".txt", ".md"]:
            with open(pfad, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        elif ext == ".pdf":
            text = ""
            try:
                import pdfplumber
                with pdfplumber.open(pfad) as pdf:
                    seiten_texte = []
                    for page in pdf.pages:
                        extracted = page.extract_text()
                        if extracted:
                            seiten_texte.append(extracted)
                    text = "\n".join(seiten_texte)
            except Exception as e:
                print(f"[Warnung] pdfplumber fehlgeschlagen bei {pfad}: {e}")

            if not text or len(text.strip()) < 50:
                try:
                    from pypdf import PdfReader
                    reader = PdfReader(pfad)
                    text = "\n".join(page.extract_text() or "" for page in reader.pages)
                except Exception as e:
                    print(f"[Warnung] pypdf-Fallback fehlgeschlagen bei {pfad}: {e}")

            if not text or len(text.strip()) < 50:
                text = ocr.pdf_text_erkennen(pfad)

            return text
        elif ext == ".docx":
            from docx import Document
            doc = Document(pfad)
            text_teile = [absatz.text for absatz in doc.paragraphs if absatz.text.strip()]

            # WICHTIG: Viele moderne Vorlagen (z.B. Lebenslauf-Layouts mit
            # Spalten für Datum/Firma/Beschreibung) speichern ihren Text in
            # TABELLEN statt in normalen Absätzen. doc.paragraphs erfasst
            # das nicht - solche Dokumente wurden bisher fälschlich als
            # "ohne Inhalt" markiert, obwohl sie voller Text waren.
            for tabelle in doc.tables:
                for zeile in tabelle.rows:
                    for zelle in zeile.cells:
                        for absatz in zelle.paragraphs:
                            if absatz.text.strip():
                                text_teile.append(absatz.text)

            return "\n".join(text_teile)
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
            text_teile = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    zeilen_text = " ".join(str(cell) for cell in row if cell is not None)
                    if zeilen_text.strip():
                        text_teile.append(zeilen_text)
            return "\n".join(text_teile)
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(pfad)
            text_teile = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_teile.append(shape.text)
            return "\n".join(text_teile)
        else:
            return None
    except PermissionError:
        # Getrennt behandelt: das ist kein defektes Dokument, sondern eine
        # fehlende Freigabe - dagegen hilft nur die Systemeinstellung.
        print(f"[Hinweis] Keine Leseberechtigung: {pfad}")
        return None
    except Exception as e:
        print(f"[Warnung] Datei konnte nicht gelesen werden: {pfad} ({e})")
        return None


# Groesse der Textabschnitte, die einzeln in einen Bedeutungsvektor
# umgerechnet werden.
#
# Warum klein: Ein Vektor ist im Kern ein Durchschnitt ueber alles, was in
# seinem Abschnitt steht. Bei 700 Zeichen einer Rechnung landen darin
# Absender, Empfaenger, Kundennummer, Zahlungsziel, Steuersatz UND die
# Positionszeile mit dem eigentlichen Artikel. Der Vektor sagt dann nur
# noch "Geschaeftsbrief mit Zahlen" - die Information, worum es
# tatsaechlich geht, verschwindet im Mittelwert.
#
# Bei rund 350 Zeichen bleibt eine Positionszeile ("1 Stk Kyocera ECOSYS
# M2540idn Multifunktionssystem s/w") in einem eigenen Abschnitt und
# behaelt ein eigenes, klares Signal. Das ist der wirksamste einzelne
# Hebel fuer die inhaltliche Suche in formularartigen Dokumenten -
# Rechnungen, Lieferscheinen, Vertraegen mit Anlagenverzeichnis.
#
# Preis: etwa doppelt so viele Abschnitte, also groesserer Index und
# laengere erste Indexierung. Das ist es wert.
ABSCHNITT_GROESSE = 350
ABSCHNITT_UEBERLAPPUNG = 80


def in_abschnitte_teilen(text, groesse=ABSCHNITT_GROESSE, ueberlappung=ABSCHNITT_UEBERLAPPUNG):
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=groesse,
            chunk_overlap=ueberlappung,
            separators=["\n\n", "\n", ".", " ", ""]
        )
        return splitter.split_text(text)
    except ImportError:
        # Notfallweg, falls langchain_text_splitters fehlt: grob nach
        # Wortzahl teilen. Rund 50 Woerter entsprechen den 350 Zeichen oben.
        woerter = text.split()
        WOERTER_PRO_ABSCHNITT = 50
        UEBERLAPPUNG_WOERTER = 12
        if len(woerter) <= WOERTER_PRO_ABSCHNITT:
            return [" ".join(woerter)]
        abschnitte = []
        schritt = max(WOERTER_PRO_ABSCHNITT - UEBERLAPPUNG_WOERTER, 1)
        for i in range(0, len(woerter), schritt):
            abschnitte.append(" ".join(woerter[i:i + WOERTER_PRO_ABSCHNITT]))
        return abschnitte


# Punkt 16 der Liste: sehr grosse Dateien. Eine einzelne 500-MB-PDF kann
# die Indexierung minutenlang blockieren und viel Arbeitsspeicher belegen -
# fuer ein Dokument, das ohnehin fast nie gesucht wird. Solche Dateien
# werden uebersprungen und in der Fehlerliste benannt.
MAX_DATEIGROESSE = 120 * 1024 * 1024  # 120 MB


# Ordner, auf die macOS den Zugriff verweigert hat. Wird bei jedem Lauf
# neu gefuellt und von der GUI ausgelesen, damit der Nutzer erfaehrt,
# warum ein Ordner leer bleibt (Punkt 5 der Liste).
verweigerte_ordner = []


def dateien_im_ordner(ordner):
    gefunden = []
    verweigerte_ordner.clear()

    def bei_fehler(fehler):
        """os.walk meldet Fehler nur ueber diesen Rueckruf - ohne ihn
        werden sie stillschweigend verschluckt."""
        if isinstance(fehler, PermissionError):
            verweigerte_ordner.append(fehler.filename or ordner)
            print(f"[Hinweis] Kein Zugriff auf {fehler.filename} - "
                  f"in den Systemeinstellungen unter Datenschutz erlauben.")
        else:
            print(f"[Warnung] Ordner nicht lesbar: {fehler}")

    for wurzel, unterordner, dateien in os.walk(ordner, onerror=bei_fehler):
        # Ignorierte Unterordner direkt aus der Traversierung entfernen
        # (verändert die Liste in-place, damit os.walk gar nicht erst
        # hineinschaut - schneller als hinterher zu filtern).
        # Safari legt einen unfertigen Download als PAKET an - einen Ordner
        # namens "bericht.pdf.download", in dem die halbe Datei liegt.
        # os.walk lief bisher hinein, fand dort ein angefangenes PDF und
        # meldete es dauerhaft als "nicht lesbar". Es ist aber schlicht ein
        # abgebrochener Download, kein Dokument des Nutzers.
        unterordner[:] = [
            u for u in unterordner
            if u not in IGNORIERTE_ORDNERNAMEN and not u.endswith(".download")
        ]
        for name in dateien:
            # Temporaere Sperrdateien von Word/Excel/PowerPoint ("~$bericht
            # .docx") und macOS-Ressourcendateien ("._datei.pdf") sind keine
            # echten Dokumente. Sie liessen sich nie lesen und landeten
            # deshalb dauerhaft in der Liste der nicht lesbaren Dateien -
            # eine Warnung, gegen die der Nutzer nichts tun kann.
            if name.startswith("~$") or name.startswith("._"):
                continue
            if not name.lower().endswith(UNTERSTUETZT):
                continue
            voller_pfad = os.path.join(wurzel, name)
            try:
                groesse = os.path.getsize(voller_pfad)
                if groesse == 0:
                    # Leere Datei: kein Fehler, nur nichts zu holen. Wurde
                    # bisher als "nicht lesbar" gemeldet und beunruhigte
                    # ohne Anlass.
                    continue
                if groesse > MAX_DATEIGROESSE:
                    print(f"[Hinweis] Uebersprungen, zu gross: {name}")
                    continue
            except OSError:
                # Groesse nicht feststellbar - dann trotzdem versuchen.
                pass
            gefunden.append(voller_pfad)
    return gefunden


def nicht_zugaengliche_ordner():
    """Ordner, die beim letzten Lauf wegen fehlender Berechtigung
    uebersprungen wurden."""
    return list(dict.fromkeys(verweigerte_ordner))

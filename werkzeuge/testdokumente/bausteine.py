"""
bausteine.py - schreibt ein Testdokument in seinem Dateiformat.

Jedes Dokument ist ein Dok(...) aus haushalt.py oder kanzlei.py. Hier steht
nur, WIE es auf die Platte kommt - fuer jedes Format, das die App liest
(siehe dateien.UNTERSTUETZT):

  pdf    PDF mit Textebene, wie aus Word oder einer Buchhaltung exportiert
  scan   PDF, das nur aus einem Bild besteht - wie vom Scanner. Leicht
         schief, verrauscht, mit Eingangsstempel. Die App kommt an den Text
         nur ueber die Texterkennung (kern/ocr.py).
  docx   Word
  xlsx   Excel
  pptx   PowerPoint
  txt/md reiner Text

GLEICHES ERGEBNIS BEI JEDEM LAUF
--------------------------------
Zufall (Schraeglage, Rauschen) kommt aus einem Zufallsgenerator, der mit
dem Dateinamen startet. Datumsangaben in den Dateien sind fest. Zip-basierte
Formate (docx, xlsx, pptx) bekommen feste Zeitstempel. So aendert ein
erneuter Lauf nur die Dateien, deren Inhalt sich wirklich geaendert hat.
"""

import datetime
import hashlib
import io
import os
import random
import re
import zipfile
from dataclasses import dataclass, field

FESTES_DATUM = datetime.datetime(2024, 9, 1, 12, 0, 0)
HINWEIS = "Frei erfundenes Testdokument fuer SmartSearch"


@dataclass
class Dok:
    """Ein Testdokument.

    datei   Pfad unter tests/testdokumente/ (Unterordner erlaubt)
    art     pdf, scan, docx, xlsx, pptx, txt oder md
    titel   nur fuer die Metadaten, taucht im Text nicht auf
    inhalt  pdf/scan/docx/txt/md: Liste von Zeilen.
              "# ..."  Ueberschrift, "## ..." Zwischenueberschrift,
              "---"    neue Seite, ""  Leerzeile
            xlsx: {Blattname: [Zeile, Zeile, ...]}, Zeile = Liste von Zellen
            pptx: [(Folientitel, [Punkt, Punkt, ...]), ...]
    stempel scan: Text des Eingangsstempels (oder None)
    """
    datei: str
    art: str
    titel: str
    inhalt: object
    stempel: str = None
    notiz: str = field(default="", repr=False)   # nur zur Doku, wird nicht geschrieben


def _zufall(dok):
    saat = int(hashlib.sha256(dok.datei.encode("utf-8")).hexdigest()[:12], 16)
    return random.Random(saat)


# ---------------------------------------------------------------------------
# PDF mit Textebene
# ---------------------------------------------------------------------------

def _pdf(dok, ziel):
    from reportlab import rl_config
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas

    rl_config.invariant = 1   # keine Zeitstempel/Zufalls-IDs im PDF
    breite, hoehe = A4
    rand = 62
    c = canvas.Canvas(ziel, pagesize=A4)
    c.setTitle(dok.titel)
    c.setSubject(HINWEIS)
    c.setAuthor("anonymous")
    y = hoehe - rand

    def neue_seite():
        nonlocal y
        c.showPage()
        y = hoehe - rand

    for zeile in dok.inhalt:
        if zeile == "---":
            neue_seite()
            continue
        if zeile.startswith("# "):
            schrift, groesse, text, abstand = "Helvetica-Bold", 14, zeile[2:], 22
        elif zeile.startswith("## "):
            schrift, groesse, text, abstand = "Helvetica-Bold", 11, zeile[3:], 17
        else:
            schrift, groesse, text, abstand = "Helvetica", 10, zeile, 13.5
        teile = simpleSplit(text, schrift, groesse, breite - 2 * rand) or [""]
        for teil in teile:
            if y < rand + abstand:
                neue_seite()
            c.setFont(schrift, groesse)
            c.drawString(rand, y, teil)
            y -= abstand
    c.save()


# ---------------------------------------------------------------------------
# Scan: ein Bild je Seite, kein Text im PDF
# ---------------------------------------------------------------------------

_SCHRIFTEN = {
    "normal": ["/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
               "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
               "C:/Windows/Fonts/times.ttf"],
    "fett": ["/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
             "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf",
             "C:/Windows/Fonts/timesbd.ttf"],
}


def _schrift(art, groesse):
    from PIL import ImageFont
    for pfad in _SCHRIFTEN[art]:
        if os.path.isfile(pfad):
            return ImageFont.truetype(pfad, groesse)
    raise SystemExit("Keine passende Schrift fuer die Scans gefunden "
                     f"(gesucht: {_SCHRIFTEN[art]}).")


def _seiten_aufteilen(zeilen):
    seiten, aktuell = [], []
    for z in zeilen:
        if z == "---":
            seiten.append(aktuell)
            aktuell = []
        else:
            aktuell.append(z)
    seiten.append(aktuell)
    return seiten


def _scan_seite(zeilen, zufall, stempel):
    """Eine A4-Seite bei 200 dpi, so wie ein Buerokopierer sie liefert."""
    import numpy as np
    from PIL import Image, ImageDraw, ImageFilter

    dpi = 200
    b, h = int(8.27 * dpi), int(11.69 * dpi)
    bild = Image.new("L", (b, h), 255)
    zeichnen = ImageDraw.Draw(bild)
    normal, fett, gross = _schrift("normal", 27), _schrift("fett", 27), _schrift("fett", 38)
    rand_x, y = 150, 170
    max_breite = b - 2 * rand_x

    for zeile in zeilen:
        if zeile.startswith("# "):
            schrift, text, abstand = gross, zeile[2:], 58
        elif zeile.startswith("## "):
            schrift, text, abstand = fett, zeile[3:], 44
        else:
            schrift, text, abstand = normal, zeile, 38
        # Spalten (zwei oder mehr Leerzeichen) wie mit Tabstopps setzen:
        # jedes Stueck beginnt an der Stelle, an der es in der Textzeile steht
        if "  " in text.strip():
            zeichenbreite = zeichnen.textlength("n", font=schrift) * 1.02
            for treffer in re.finditer(r"\S+(?: \S+)*", text):
                x = rand_x + treffer.start() * zeichenbreite
                zeichnen.text((x, y), treffer.group(), font=schrift, fill=zufall.randint(20, 45))
            y += abstand
            continue
        # Zeilen umbrechen, die nicht auf die Seite passen
        woerter, teil = text.split(" "), ""
        stuecke = []
        for w in woerter:
            probe = (teil + " " + w).strip()
            if zeichnen.textlength(probe, font=schrift) > max_breite and teil:
                stuecke.append(teil)
                teil = w
            else:
                teil = probe
        stuecke.append(teil)
        for s in stuecke:
            zeichnen.text((rand_x, y), s, font=schrift, fill=zufall.randint(20, 45))
            y += abstand

    if stempel:
        # Eingangsstempel schraeg oben rechts, blasser Rahmen
        st = Image.new("L", (560, 150), 255)
        sz = ImageDraw.Draw(st)
        sz.rectangle([4, 4, 555, 145], outline=110, width=5)
        sz.text((30, 22), stempel.split("|")[0], font=_schrift("fett", 34), fill=110)
        if "|" in stempel:
            sz.text((30, 82), stempel.split("|")[1], font=_schrift("normal", 30), fill=110)
        st = st.rotate(zufall.uniform(-9, -4), expand=True, fillcolor=255)
        bild.paste(st, (b - st.width - 110, 60),
                   mask=Image.eval(st, lambda v: 255 - v))

    # Schraeg eingelegt, leicht unscharf, Grundrauschen, grauer Hintergrund
    bild = bild.rotate(zufall.uniform(-1.1, 1.1), resample=Image.BICUBIC, fillcolor=250)
    bild = bild.filter(ImageFilter.GaussianBlur(0.6))
    feld = np.asarray(bild, dtype="float32")
    rng = np.random.default_rng(zufall.randint(0, 2**31))
    feld = feld * 0.93 + 8 + rng.normal(0, 7, feld.shape)
    # ein paar Staubpunkte
    for _ in range(40):
        x, yy = rng.integers(0, b), rng.integers(0, h)
        feld[max(0, yy - 1):yy + 2, max(0, x - 1):x + 2] = rng.integers(60, 140)
    return Image.fromarray(np.clip(feld, 0, 255).astype("uint8"))


def _scan(dok, ziel):
    from reportlab import rl_config
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    rl_config.invariant = 1
    zufall = _zufall(dok)
    c = canvas.Canvas(ziel, pagesize=A4)
    # Scanner schreiben keinen Titel - nur den Geraetenamen
    c.setTitle("")
    c.setSubject(HINWEIS)
    c.setCreator("Scan to PDF")
    for nr, seite in enumerate(_seiten_aufteilen(dok.inhalt)):
        bild = _scan_seite(seite, zufall, dok.stempel if nr == 0 else None)
        puffer = io.BytesIO()
        bild.save(puffer, "JPEG", quality=72, optimize=True)
        puffer.seek(0)
        c.drawImage(ImageReader(puffer), 0, 0, *A4)
        c.showPage()
    c.save()


# ---------------------------------------------------------------------------
# Office-Formate
# ---------------------------------------------------------------------------

def _docx(dok, ziel):
    from docx import Document
    from docx.enum.text import WD_BREAK

    d = Document()
    d.core_properties.title = dok.titel
    d.core_properties.subject = HINWEIS
    d.core_properties.author = "anonymous"
    d.core_properties.last_modified_by = "anonymous"
    d.core_properties.created = FESTES_DATUM
    d.core_properties.modified = FESTES_DATUM
    d.core_properties.revision = 1
    for zeile in dok.inhalt:
        if zeile == "---":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif zeile.startswith("# "):
            d.add_heading(zeile[2:], level=1)
        elif zeile.startswith("## "):
            d.add_heading(zeile[3:], level=2)
        elif zeile:
            d.add_paragraph(zeile)
    d.save(ziel)


def _xlsx(dok, ziel):
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.title = dok.titel
    wb.properties.subject = HINWEIS
    wb.properties.creator = "anonymous"
    wb.properties.created = FESTES_DATUM
    wb.properties.modified = FESTES_DATUM
    for blatt, zeilen in dok.inhalt.items():
        ws = wb.create_sheet(blatt)
        for zeile in zeilen:
            ws.append(zeile)
        for zelle in ws[1]:
            zelle.font = Font(bold=True)
        for spalte in ws.columns:
            laenge = max(len(str(z.value)) for z in spalte if z.value is not None)
            ws.column_dimensions[spalte[0].column_letter].width = min(60, laenge + 2)
    wb.save(ziel)


def _pptx(dok, ziel):
    from pptx import Presentation

    p = Presentation()
    p.core_properties.title = dok.titel
    p.core_properties.subject = HINWEIS
    p.core_properties.author = "anonymous"
    p.core_properties.last_modified_by = "anonymous"
    p.core_properties.created = FESTES_DATUM
    p.core_properties.modified = FESTES_DATUM
    p.core_properties.revision = 1
    for nr, (titel, punkte) in enumerate(dok.inhalt):
        if nr == 0:
            folie = p.slides.add_slide(p.slide_layouts[0])
            folie.shapes.title.text = titel
            folie.placeholders[1].text = "\n".join(punkte)
        else:
            folie = p.slides.add_slide(p.slide_layouts[1])
            folie.shapes.title.text = titel
            rahmen = folie.placeholders[1].text_frame
            rahmen.text = punkte[0] if punkte else ""
            for punkt in punkte[1:]:
                rahmen.add_paragraph().text = punkt
    p.save(ziel)


def _zip_festschreiben(pfad):
    """Feste Zeitstempel in docx/xlsx/pptx (das sind Zip-Archive)."""
    with zipfile.ZipFile(pfad) as z:
        teile = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(pfad, "w", zipfile.ZIP_DEFLATED) as z:
        for name, daten in teile:
            if name == "docProps/core.xml":
                # openpyxl setzt "zuletzt geaendert" beim Speichern immer auf jetzt
                daten = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*",
                               rb"\g<1>" + FESTES_DATUM.strftime("%Y-%m-%dT%H:%M:%SZ").encode(), daten)
            info = zipfile.ZipInfo(name, date_time=FESTES_DATUM.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, daten)


def _text(dok, ziel):
    zeilen = []
    for z in dok.inhalt:
        if z == "---":
            continue
        if dok.art == "txt" and z.startswith("#"):
            z = z.lstrip("# ").upper()
        zeilen.append(z)
    with open(ziel, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(zeilen) + "\n")


SCHREIBER = {"pdf": _pdf, "scan": _scan, "docx": _docx, "xlsx": _xlsx,
             "pptx": _pptx, "txt": _text, "md": _text}


def schreiben(dok, basis):
    ziel = os.path.join(basis, dok.datei)
    os.makedirs(os.path.dirname(ziel), exist_ok=True)
    SCHREIBER[dok.art](dok, ziel)
    if dok.art in ("docx", "xlsx", "pptx"):
        _zip_festschreiben(ziel)
    # Aenderungsdatum fest - die App kann nach Zeitraum filtern
    zeit = FESTES_DATUM.timestamp()
    os.utime(ziel, (zeit, zeit))
    return ziel
